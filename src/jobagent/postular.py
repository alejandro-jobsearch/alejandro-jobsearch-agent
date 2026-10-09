"""Local application step (human-in-the-loop): tracked Issue → tailored CV .docx → ATS after → Issue/board.

    python -m jobagent.postular fetch   --data-dir D --issue 1 --work work/1
    python -m jobagent.postular build   --data-dir D --work work/1 --plan work/1/plan.yml --out "CV.docx" [--pages]
    python -m jobagent.postular publish --data-dir D --issue 1 --work work/1
    python -m jobagent.postular applied --data-dir D --issue 1 --work work/1     # after the human sends it

The writing (which verified achievements to use and how to phrase them for this posting) is done by the
operator — in practice Claude Code following the `postular` skill — into `plan.yml`:

    lang: en
    headline: "..."
    summary: "..."
    roles:                       # keyed by the role title exactly as it appears in the master CV
      Cloud Solutions Architect: ["bullet", "bullet"]
    skills:                      # optional: replace a skills line by its label prefix
      "Cloud & FinOps": "Cloud & FinOps: ..."
    sources: {"Cloud Solutions Architect": ["csa-ai-arch", ...]}   # achievement ids backing each role (audit)

`build` copies the master CV, swaps only text (formatting, styles and layout are kept), then scores ATS
coverage against the posting's keywords and checks the page count with Word when available.
"""
import argparse
import copy
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import docx
import yaml

from jobagent import ats

SUMMARY_HEADINGS = {"PROFESSIONAL SUMMARY", "RESUMEN PROFESIONAL", "PERFIL PROFESIONAL"}
SECTION_HEADINGS = {"EDUCATION", "EDUCACIÓN", "FORMACIÓN ACADÉMICA", "SKILLS & CERTIFICATIONS",
                    "COMPETENCIAS & CERTIFICACIONES", "WORK EXPERIENCE", "EXPERIENCIA PROFESIONAL", "EXPERIENCIA LABORAL"}


# --- GitHub ------------------------------------------------------------------------------------
def _tracker(data_dir: Path):
    from jobagent.tracker import Tracker

    cfg = yaml.safe_load((data_dir / "config" / "search.yml").read_text(encoding="utf-8"))["tracker"]
    token = os.environ.get("GH_PROJECT_TOKEN") or subprocess.run(
        ["gh", "auth", "token"], capture_output=True, text=True, check=True).stdout.strip()
    repo = os.environ.get("GITHUB_REPOSITORY") or f'{cfg["org"]}/{data_dir.resolve().name}'
    return Tracker(token, repo, cfg["org"], cfg["project_number"])


def fetch(data_dir: Path, issue: int, work: Path) -> dict:
    body = _tracker(data_dir).issue(issue)["body"] or ""
    m = re.search(r"<!-- jobagent-meta: (\{.*?\}) -->", body, re.S)
    if not m:
        sys.exit(f"Issue #{issue} has no jobagent-meta block (created before phase 5?)")
    meta = json.loads(m.group(1))
    d = re.search(r"<summary>Descripción original</summary>\s*(.*?)</details>", body, re.S)
    meta["description"] = d.group(1).strip() if d else ""
    work.mkdir(parents=True, exist_ok=True)
    (work / "job.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"#{issue} {meta['company']} — {meta['title']} · lang={meta['lang']} · score={meta['score']} · "
          f"ATS antes={meta['ats_before']}% · {len(meta['keywords'])} keywords → {work / 'job.json'}")
    return meta


# --- docx --------------------------------------------------------------------------------------
def _set_text(p, text: str) -> None:
    """Replace a paragraph's text keeping the first run's formatting."""
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def _is_bullet(p) -> bool:
    return p.style.name.startswith("List")


def apply_plan(master: Path, plan: dict, out: Path) -> None:
    d = docx.Document(master)
    ps = d.paragraphs
    if plan.get("headline"):
        _set_text(ps[1], plan["headline"])
    if plan.get("summary"):
        i = next(i for i, p in enumerate(ps) if p.text.strip().upper() in SUMMARY_HEADINGS)
        _set_text(ps[i + 1], plan["summary"])

    for title, bullets in (plan.get("roles") or {}).items():
        ps = d.paragraphs
        try:
            t = next(i for i, p in enumerate(ps) if p.text.strip() == title and not _is_bullet(p))
        except StopIteration:
            sys.exit(f"role title not found in master CV: {title!r}")
        start = t + 2                                   # title → dates → bullets
        end = start
        while end < len(ps) and _is_bullet(ps[end]):
            end += 1
        existing = ps[start:end]
        if not existing:
            sys.exit(f"no bullets under {title!r} in the master CV")
        template = copy.deepcopy(existing[0]._p)
        for p in existing[len(bullets):]:               # drop extra bullets
            p._p.getparent().remove(p._p)
        anchor = existing[min(len(bullets), len(existing)) - 1]._p
        for p, text in zip(existing, bullets):          # rewrite kept ones
            _set_text(p, text)
        for text in bullets[len(existing):]:            # add missing ones after the last kept bullet
            new = copy.deepcopy(template)
            anchor.addnext(new)
            anchor = new
            from docx.text.paragraph import Paragraph
            _set_text(Paragraph(new, existing[0]._parent), text)

    for prefix, text in (plan.get("skills") or {}).items():
        p = next((p for p in d.paragraphs if p.text.strip().startswith(prefix)), None)
        if p is None:
            sys.exit(f"skills line not found: {prefix!r}")
        _set_text(p, text)
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)


def docx_text(path: Path) -> str:
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def page_count(path: Path) -> int | None:
    """Word's own pagination (Windows only); None when Word is not available."""
    if os.name != "nt":
        return None
    ps = ("$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
          f"$d=$w.Documents.Open('{path.resolve()}',$false,$true);$d.ComputeStatistics(2);$d.Close($false);$w.Quit()")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=120)
        return int(r.stdout.strip().splitlines()[-1])
    except Exception:
        return None


def build(data_dir: Path, work: Path, plan_path: Path, out: Path, pages: bool, master_dir: Path) -> dict:
    job = json.loads((work / "job.json").read_text(encoding="utf-8"))
    plan = yaml.safe_load(plan_path.read_text(encoding="utf-8"))
    lang = plan.get("lang") or job["lang"] or "en"
    master = next(master_dir.glob(f"CV - * - {lang.upper()} 20??.docx"))
    apply_plan(master, plan, out)
    after, missing = ats.coverage(job["keywords"], docx_text(out))
    base, _ = ats.coverage(job["keywords"], docx_text(master))
    result = {"cv": str(out), "master": master.name, "lang": lang, "ats_master": base, "ats_after": after,
              "ats_before_issue": job["ats_before"], "missing": missing,
              "pages": page_count(out) if pages else None, "sources": plan.get("sources", {})}
    (work / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "sources"}, ensure_ascii=False))
    return result


def publish(data_dir: Path, issue: int, work: Path, deck: str | None) -> None:
    res = json.loads((work / "result.json").read_text(encoding="utf-8"))
    t = _tracker(data_dir)
    missing = ", ".join(res["missing"]) or "(ninguna)"
    srcs = "\n".join(f"- {role}: {', '.join(ids)}" for role, ids in res["sources"].items()) or "- (no registrado)"
    body = f"""### ✅ CV listo
**ATS:** {res['ats_before_issue']}% → **{res['ats_after']}%** · páginas: {res['pages'] or '?'} · idioma: {res['lang']}
**Archivo:** `{Path(res['cv']).name}`
{f"**Deck de preparación:** `{deck}`" if deck else ""}

**Keywords que siguen faltando** (no se agregan si no hay experiencia real que las respalde): {missing}

<details><summary>Logros del perfil usados (auditoría)</summary>

{srcs}
</details>"""
    t.comment(issue, body)
    t.set_fields(issue, ATS_despues=res["ats_after"], Status="CV listo")
    print(f"#{issue}: comentario publicado; tablero → CV listo, ATS despues={res['ats_after']}")


def applied(data_dir: Path, issue: int, work: Path, when: str) -> None:
    """After the human sends the application: history entry + board card to Postulada."""
    job = json.loads((work / "job.json").read_text(encoding="utf-8"))
    res_file = work / "result.json"
    res = json.loads(res_file.read_text(encoding="utf-8")) if res_file.exists() else {}
    ats_note = f"; CV {res['lang'].upper()} ATS {res['ats_before_issue']}→{res['ats_after']}" if res else ""
    entry = {"company": job["company"], "role": job["title"],
             "note": f"POSTULADO {when} vía pipeline (Issue #{issue}){ats_note}"}
    path = data_dir / "config" / "applied.yml"
    with path.open("a", encoding="utf-8") as f:
        f.write(yaml.safe_dump([entry], allow_unicode=True, sort_keys=False, width=200))
    t = _tracker(data_dir)
    t.set_fields(issue, Status="Postulada")
    t.comment(issue, f"📨 Postulada el {when}.")
    print(f"#{issue}: applied.yml actualizado; tablero → Postulada")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    for name in ("fetch", "build", "publish", "applied"):
        s = sub.add_parser(name)
        s.add_argument("--data-dir", required=True, type=Path)
        s.add_argument("--work", required=True, type=Path)
        if name in ("fetch", "publish", "applied"):
            s.add_argument("--issue", required=True, type=int)
        if name == "applied":
            s.add_argument("--date", default=__import__("datetime").date.today().isoformat())
        if name == "build":
            s.add_argument("--plan", required=True, type=Path)
            s.add_argument("--out", required=True, type=Path)
            s.add_argument("--master-dir", type=Path, default=Path(os.environ.get("CV_MASTER_DIR", ".")))
            s.add_argument("--pages", action="store_true", help="check page count with Word (Windows)")
        if name == "publish":
            s.add_argument("--deck", help="path of the prep deck, if one was generated")
    a = p.parse_args()
    if a.cmd == "fetch":
        fetch(a.data_dir, a.issue, a.work)
    elif a.cmd == "build":
        build(a.data_dir, a.work, a.plan, a.out, a.pages, a.master_dir)
    elif a.cmd == "publish":
        publish(a.data_dir, a.issue, a.work, a.deck)
    else:
        applied(a.data_dir, a.issue, a.work, a.date)
    return 0


if __name__ == "__main__":
    sys.exit(main())
